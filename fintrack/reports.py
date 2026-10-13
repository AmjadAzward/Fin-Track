"""PDF/XLSX export for selected month and currency."""
from datetime import datetime
from pathlib import Path
from xml.sax.saxutils import escape
from openpyxl import Workbook
from openpyxl.styles import PatternFill,Font,Alignment
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import SimpleDocTemplate,Paragraph,Spacer,Table,TableStyle


def money(cents,currency='LKR'):
    return f'{currency} {cents/100:,.2f}'


def xlsx(db,path,month,currency):
    wb=Workbook(); ws=wb.active; ws.title='Summary'
    summ=db.monthly_summary(month,currency)
    for row in [('FinTrack Finance Summary',),('Period',month),('Currency',currency),(),('Income',summ['income']/100),('Expenses',summ['expenses']/100),('Net Cash Flow',(summ['income']-summ['expenses'])/100)]:ws.append(row)
    ws['A1'].font=Font(bold=True,size=16,color='135D76')
    for col in ['B5','B6','B7']: ws[col].number_format='#,##0.00;[Red](#,##0.00)'
    tx=wb.create_sheet('Transactions');tx.append(['Date','Account','Description','Category','Amount','Currency','Reference / Note'])
    for item in db.transactions(month=month,limit=100000):
        if item['currency']==currency:tx.append([item['date'],item['account_name'],item['description'],item['category_name'] or 'Uncategorized',item['amount_cents']/100,currency,item['note']])
    for cell in tx[1]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='183149')
    for row in tx.iter_rows(min_row=2):row[4].number_format='#,##0.00;[Red](#,##0.00)'
    for sheet in wb:
        sheet.freeze_panes='A2'
        for col,w in {'A':22,'B':25,'C':42,'D':22,'E':18,'F':16,'G':28}.items():sheet.column_dimensions[col].width=w
    wb.save(path)


def pdf(db,path,month,currency):
    styles=getSampleStyleSheet();summ=db.monthly_summary(month,currency)
    doc=SimpleDocTemplate(str(path),pagesize=A4,rightMargin=34,leftMargin=34,topMargin=35,bottomMargin=35)
    items=[Paragraph('FinTrack | Monthly Finance Report',styles['Title']),Spacer(1,12),
           Paragraph(f'Period: {escape(month)} &nbsp;&nbsp; Currency: {escape(currency)} &nbsp;&nbsp; Generated: {datetime.now():%Y-%m-%d}',styles['Normal']),Spacer(1,18)]
    summary=[['Income',money(summ['income'],currency)],['Expenses',money(summ['expenses'],currency)],['Net Cash Flow',money(summ['income']-summ['expenses'],currency)]]
    t=Table(summary,colWidths=[220,200]);t.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,-1),colors.HexColor('#edf3f7')),('BOTTOMPADDING',(0,0),(-1,-1),11),('TOPPADDING',(0,0),(-1,-1),11)]));items.extend([t,Spacer(1,20),Paragraph('Transactions',styles['Heading2'])])
    data=[['Date','Account','Description','Amount']]
    for r in db.transactions(month=month,limit=100000):
        if r['currency']==currency:
            data.append([r['date'],r['account_name'][:18],Paragraph(escape(r['description'][:46]),styles['BodyText']),money(r['amount_cents'],currency)])
    if len(data)==1:data.append(['—','—','No transactions','—'])
    table=Table(data,colWidths=[72,105,200,138],repeatRows=1,hAlign='LEFT')
    table.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#183149')),('TEXTCOLOR',(0,0),(-1,0),colors.white),('FONTSIZE',(0,0),(-1,-1),8),('VALIGN',(0,0),(-1,-1),'TOP'),('ROWBACKGROUNDS',(0,1),(-1,-1),[colors.white,colors.HexColor('#f2f6fa')]),('BOTTOMPADDING',(0,0),(-1,-1),8),('TOPPADDING',(0,0),(-1,-1),8)]))
    items.append(table);doc.build(items)
