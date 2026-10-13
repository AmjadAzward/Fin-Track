"""User-controlled CSV/XLSX import. No network or automatic bank connections."""
import csv
from pathlib import Path
from datetime import datetime, date
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from openpyxl import load_workbook


def read_rows(path):
    path=Path(path)
    if path.suffix.lower()=='.csv':
        with path.open('r',encoding='utf-8-sig',newline='') as f:
            sample=f.read(4096); f.seek(0)
            try: dialect=csv.Sniffer().sniff(sample,delimiters=',;\t|')
            except csv.Error: dialect=csv.excel
            rows=list(csv.reader(f,dialect))
    elif path.suffix.lower()=='.xlsx':
        wb=load_workbook(path,read_only=True,data_only=True)
        try: rows=[list(row) for row in wb.active.iter_rows(values_only=True)]
        finally: wb.close()
    else: raise ValueError('Only CSV and XLSX files are supported in this release')
    rows=[[str(v).strip() if v is not None else '' for v in row] for row in rows]
    return [r for r in rows if any(r)]


def cents(value):
    s=str(value).strip().replace(',','').replace('LKR','').replace('Rs.','').replace('$','').strip()
    if s.startswith('(') and s.endswith(')'): s='-'+s[1:-1]
    try: return int((Decimal(s)*100).quantize(Decimal('1'),rounding=ROUND_HALF_UP))
    except (InvalidOperation,ValueError): raise ValueError(f'Invalid amount: {value}')


def parse_date(value):
    if isinstance(value,(datetime,date)): return value.strftime('%Y-%m-%d')
    s=str(value).strip()
    for fmt in ('%Y-%m-%d','%d/%m/%Y','%m/%d/%Y','%d-%m-%Y','%d.%m.%Y','%Y/%m/%d','%Y-%m-%d %H:%M:%S'):
        try:return datetime.strptime(s,fmt).date().isoformat()
        except ValueError:pass
    raise ValueError(f'Invalid date: {s}. Select a supported format.')


def convert(rows,columns,header=True):
    """columns maps date, description, amount OR debit+credit, reference to integer indices."""
    result=[]; start=1 if header else 0
    for number,row in enumerate(rows[start:],start+1):
        try:
            def get(k):
                ix=columns.get(k,-1)
                return row[ix] if isinstance(ix,int) and 0<=ix<len(row) else ''
            day=parse_date(get('date'))
            desc=get('description').strip()
            if 'amount' in columns:
                amount=cents(get('amount'))
            else:
                debit=cents(get('debit')) if get('debit') else 0
                credit=cents(get('credit')) if get('credit') else 0
                amount=credit-debit
            if not desc or not amount: raise ValueError('Description and non-zero amount are required')
            result.append({'date':day,'description':desc,'amount_cents':amount,'reference':get('reference')})
        except (ValueError,IndexError) as exc:
            raise ValueError(f'Row {number}: {exc}') from exc
    return result
