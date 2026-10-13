import csv
import pytest
from fintrack.db import FinanceDB


def test_edit_transaction_and_category_validation(tmp_path):
    db=FinanceDB(tmp_path/'edit.db')
    a=db.add_account('Main')
    cats={r['name']:r['id'] for r in db.categories()}
    ident=db.add_transaction(a,'2026-10-01','Original',-1000,cats['Groceries'])
    db.edit_transaction(ident,'2026-10-02','Updated',-2500,cats['Transport'],'changed')
    r=db.transactions()[0]
    assert r['description']=='Updated' and r['amount_cents']==-2500 and r['note']=='changed'
    assert db.monthly_summary('2026-10')['expenses']==2500
    with pytest.raises(ValueError):
        db.edit_transaction(ident,'2026-10-02','Wrong type',-200,cats['Salary'])
    assert db.database_check()=='ok'
    db.close()


def test_import_renamed_files_and_auto_category(tmp_path):
    db=FinanceDB(tmp_path/'import.db')
    a=db.add_account('Checking')
    rows=[{'date':'2026-10-01','description':'Supermarket groceries','amount_cents':-1750},
          {'date':'2026-10-01','description':'Salary credit','amount_cents':50000}]
    assert db.import_records(a,'a.csv',rows)==(2,0)
    assert db.import_records(a,'renamed.csv',rows)==(0,2)
    cats={x['description']:x['category_name'] for x in db.transactions()}
    assert cats['Supermarket groceries']=='Groceries'
    assert cats['Salary credit']=='Salary'
    assert db.monthly_summary('2026-10')['income']==50000
    assert db.monthly_summary('2026-10')['expenses']==1750
    db.close()


def test_csv_export_and_integrity(tmp_path):
    db=FinanceDB(tmp_path/'data.db')
    acc=db.add_account('Main')
    db.add_transaction(acc,'2026-10-01','Lunch',-1299)
    path=tmp_path/'report.csv'
    db.export_transactions_csv(path,'2026-10','LKR')
    with path.open(encoding='utf-8-sig',newline='') as f:
        rows=list(csv.reader(f))
    assert rows[1][2]=='Lunch' and rows[1][4]=='-12.99'
    assert db.database_check()=='ok'
    db.close()


def test_invalid_transfer_date(tmp_path):
    db=FinanceDB(tmp_path/'data.db')
    a=db.add_account('A');b=db.add_account('B')
    with pytest.raises(ValueError):db.transfer(a,b,'invalid',5000)
    assert db.transactions()==[]
    db.close()


def test_pdf_and_excel_reports(tmp_path):
    from fintrack.reports import pdf,xlsx
    from openpyxl import load_workbook
    db=FinanceDB(tmp_path/'reports.db')
    a=db.add_account('Bank')
    cats={r['name']:r['id'] for r in db.categories()}
    db.add_transaction(a,'2026-10-11','Freelance',250000,cats['Freelance'])
    db.add_transaction(a,'2026-10-11','Groceries',-9500,cats['Groceries'])
    pdf_file=tmp_path/'statement.pdf';xlsx_file=tmp_path/'statement.xlsx'
    pdf(db,pdf_file,'2026-10','LKR');xlsx(db,xlsx_file,'2026-10','LKR')
    assert pdf_file.read_bytes().startswith(b'%PDF')
    wb=load_workbook(xlsx_file,data_only=True)
    assert wb['Summary']['B5'].value==2500
    assert wb['Summary']['B6'].value==95
    assert wb['Transactions'].max_row==3
    wb.close();db.close()
